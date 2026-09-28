from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO
from datetime import datetime

def export_chat_to_docx(messages: list[dict], title: str = "Báo cáo Hỏi Đáp Học Thuật") -> bytes:
    doc = DocxDocument()
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.5)

    h = doc.add_heading(title, 0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info = doc.add_paragraph()
    info.add_run(f"Ngày xuất: {datetime.now().strftime('%d/%m/%Y %H:%M')}").italic = True
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("─" * 60)

    q_count = 0
    for msg in messages:
        if not msg.get("content", "").strip():
            continue
        if msg["role"] == "user":
            q_count += 1
            p = doc.add_paragraph()
            run = p.add_run(f"Câu hỏi {q_count}: {msg['content']}")
            run.bold = True
            run.font.color.rgb = RGBColor(0, 70, 180)
            run.font.size = Pt(12)
        else:
            p = doc.add_paragraph()
            run = p.add_run(msg["content"].replace("**", "").replace("*", ""))
            run.font.size = Pt(11)
            p.paragraph_format.left_indent = Cm(0.5)
        doc.add_paragraph("")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
